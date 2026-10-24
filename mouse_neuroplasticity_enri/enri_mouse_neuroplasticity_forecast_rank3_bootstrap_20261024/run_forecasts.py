#!/usr/bin/env python3
"""Bootstrap aggregation (bagging) of regressors fitted on tiny real-mouse cohorts.

Inside each leave-one-mouse-out fold, resample ONLY the remaining real mice
with replacement, refit, and average predictions. No pseudo-independent mice.
Training scaling fixed on the outer training partition; no tuning on test mice.
"""
from pathlib import Path
import json
import warnings
import numpy as np
import pandas as pd
from sklearn.linear_model import Ridge, BayesianRidge, HuberRegressor
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Patch, Rectangle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
OUT=HERE/'results';OUT.mkdir(parents=True,exist_ok=True)
DATA=HERE/'input'/'clean133_k3_rank_frozen_scores.csv'
DOMAINS=('D1','D2','D3')
INTERVALS=[('PBS 0→16','PBS',0,16,9),('PBS 16→24','PBS',16,24,6),('LPS 16→24','LPS',16,24,11),('PBS 0→24','PBS',0,24,6)]
ALGOS=['Ridge','Bayesian Ridge','Huber']
REPS=100; CHECK_REPS=(25,50,100); SEED=20261024
GROUP_COLORS={'PBS':'#527F9D','LPS':'#B87959'}
COLORS={'Ridge':'#387B9E','Bayesian Ridge':'#6390B1','Huber':'#56917C','Без изменений':'#CE9571','Среднее':'#94A6B0'}
INK='#2E4C5D';GRAY='#637985'


def pairs(df,group,t0,t1):
    a=df[(df.group==group)&(df.week==t0)].copy()
    b=df[(df.group==group)&(df.week==t1)].copy()
    assert a.mouse_id.is_unique and b.mouse_id.is_unique
    x=a[['mouse_id','status']+[d+'_score' for d in DOMAINS]].rename(columns={'status':'start_status',**{d+'_score':d+'_start' for d in DOMAINS}})
    y=b[['mouse_id','status']+[d+'_score' for d in DOMAINS]].rename(columns={'status':'end_status',**{d+'_score':d+'_end' for d in DOMAINS}})
    return x.merge(y,on='mouse_id',validate='one_to_one').sort_values('mouse_id').reset_index(drop=True)


def model(label):
    if label=='Ridge':return Ridge(alpha=1.,fit_intercept=True)
    if label=='Bayesian Ridge':return BayesianRidge(max_iter=300,tol=1e-3,alpha_1=1e-6,alpha_2=1e-6,lambda_1=1e-6,lambda_2=1e-6,fit_intercept=True)
    if label=='Huber':return HuberRegressor(alpha=.1,epsilon=1.35,max_iter=1500,tol=1e-6,fit_intercept=True)
    raise ValueError(label)


def train_predict(z,y,z_query,algorithm):
    if np.std(z)<1.e-12:
        return float(np.mean(y))
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        m=model(algorithm).fit(z.reshape(-1,1),y)
        return float(m.predict(np.array([[z_query]]))[0])


def fold(x,y,heldout,task_seed):
    train=np.arange(len(x))!=heldout
    rawx=x[train].astype(float); target=y[train].astype(float)
    mu=rawx.mean();sd=rawx.std(ddof=0)
    if sd<1.e-12:
        z=np.zeros_like(rawx);q=0.
    else:
        z=(rawx-mu)/sd;q=(float(x[heldout])-mu)/sd
    rng=np.random.default_rng(task_seed)
    indices=rng.integers(0,len(z),size=(REPS,len(z)))
    unique_counts=np.array([len(np.unique(ix)) for ix in indices])
    out={}
    for alg in ALGOS:
        direct=train_predict(z,target,q,alg)
        fitted=np.array([train_predict(z[ix],target[ix],q,alg) for ix in indices])
        out[f'{alg}|direct']=direct
        for rep in CHECK_REPS:
            out[f'{alg}|bootstrap{rep}']=float(fitted[:rep].mean())
        out[f'{alg}|bootstrap_sd']=float(fitted.std(ddof=1))
    out['Без изменений']=float(x[heldout]);out['Среднее']=float(target.mean())
    out['mean_unique_bootstrap_training_mice']=float(unique_counts.mean())
    return out


def calculate(df):
    metrics=[]; predrows=[]; paired=[]
    allfolds=0
    for taskidx,(title,group,t0,t1,expected) in enumerate(INTERVALS):
        p=pairs(df,group,t0,t1)
        assert len(p)==expected,(title,len(p))
        pp=p.copy();pp.insert(0,'transition',title);paired.append(pp)
        for domain_idx,domain in enumerate(DOMAINS):
            x=p[domain+'_start'].to_numpy(float);y=p[domain+'_end'].to_numpy(float)
            foldout=[]
            for heldout in range(len(x)):
                out=fold(x,y,heldout,SEED+taskidx*1000000+domain_idx*100000+heldout*1000)
                foldout.append(out)
                predrows.append({'transition':title,'group':group,'domain':domain,'mouse_id':p.mouse_id.iloc[heldout],
                        'status_start':p.start_status.iloc[heldout],'status_end':p.end_status.iloc[heldout],
                        'current':float(x[heldout]),'actual':float(y[heldout]),**out})
                allfolds+=1
            frame=pd.DataFrame(foldout)
            mae=lambda key:float(np.abs(y-frame[key].to_numpy(float)).mean())
            rmse=lambda key:float(np.sqrt(np.mean((y-frame[key].to_numpy(float))**2)))
            stats={'transition':title,'group':group,'week_from':t0,'week_to':t1,'domain':domain,'n':len(x),
                   'mean_unique_training_mice':float(frame['mean_unique_bootstrap_training_mice'].mean())}
            for alg in ALGOS:
                slug={'Ridge':'ridge','Bayesian Ridge':'bayes','Huber':'huber'}[alg]
                for key,repr_ in [('direct','plain'),('bootstrap25','boot25'),('bootstrap50','boot50'),('bootstrap100','boot100')]:
                    stats[f'{slug}_{repr_}_mae']=mae(f'{alg}|{key}')
                    stats[f'{slug}_{repr_}_rmse']=rmse(f'{alg}|{key}')
                stats[f'{slug}_boot_gain']=stats[f'{slug}_plain_mae']-stats[f'{slug}_boot100_mae']
            for base in ['Без изменений','Среднее']:
                name='unchanged' if base=='Без изменений' else 'mean'
                stats[name+'_mae']=mae(base);stats[name+'_rmse']=rmse(base)
            candidates={k:mae(k) for k in ['Без изменений','Среднее']+[f'{a}|{p_}' for a in ALGOS for p_ in ['direct','bootstrap100']]}
            stats['best_method']=min(candidates,key=candidates.get)
            stats['best_mae']=min(candidates.values())
            stats['bootstrap_best_of_8']=stats['best_method'].endswith('|bootstrap100')
            metrics.append(stats)
    m=pd.DataFrame(metrics);p=pd.DataFrame(predrows)
    assert len(m)==12 and len(p)==96,(len(m),len(p))
    # Recalculate historical 12 direct Ridge MAEs precisely to verify comparability.
    old={('PBS 0→16','D1'):1.2000215528785045,('PBS 0→16','D2'):.881626261705343,('PBS 0→16','D3'):1.009777814049385,
         ('PBS 16→24','D1'):.27062259785156545,('PBS 16→24','D2'):.9565559943169326,('PBS 16→24','D3'):1.610202511258489,
         ('LPS 16→24','D1'):.7244930108270334,('LPS 16→24','D2'):.3224514013325677,('LPS 16→24','D3'):1.1515697700575693,
         ('PBS 0→24','D1'):.2906802250739176,('PBS 0→24','D2'):.6213566159962042,('PBS 0→24','D3'):.770379377664982}
    assert all(abs(r.ridge_plain_mae-old[(r.transition,r.domain)])<.000001 for r in m.itertuples()),'Ridge mismatch'
    m.to_csv(OUT/'model_metrics.csv',index=False)
    p.to_csv(OUT/'leave_one_mouse_out_predictions.csv',index=False)
    pd.concat(paired,ignore_index=True).to_csv(OUT/'paired_mouse_scores.csv',index=False)
    return m,p

def spread(vals,lo,hi,min_gap):
    a=sorted(vals,key=lambda t:(t[1],t[0]))
    v=[max(lo,y) for _,y in a]
    for i in range(1,len(v)):v[i]=max(v[i],v[i-1]+min_gap)
    if v and v[-1]>hi:v=[y-(v[-1]-hi) for y in v]
    for i in range(len(v)-2,-1,-1):v[i]=min(v[i],v[i+1]-min_gap)
    if v and v[0]<lo:v=[y+lo-v[0] for y in v]
    return {a[i][0]:v[i] for i in range(len(a))}


def plot_traj(df,domain):
    fig,axes=plt.subplots(1,2,figsize=(15.8,3.9),dpi=155,sharey=True)
    vals=df[df.group.isin(['PBS','LPS'])][f'{domain}_score']
    lo=float(vals.min())-.45;hi=float(vals.max())+.45
    for ax,group in zip(axes,['PBS','LPS']):
        subset=df[df.group==group];ends={}
        for id_,g in subset.groupby('mouse_id'):
            g=g.sort_values('week');xx=g.week.to_numpy();yy=g[f'{domain}_score'].to_numpy()
            ax.plot(xx,yy,lw=1.3,alpha=.65,color=GROUP_COLORS[group],marker='o',ms=3.5,zorder=2)
            ends[id_]=(xx[-1],yy[-1])
        for time in (0,16,24):ax.axvline(time,color='#E3E8ED',lw=.6,zorder=0)
        ax.axhline(0,color='#ADB9C1',ls=':',lw=.8)
        for time in (0,16,24):
            e=[(name,y) for name,(x,y) in ends.items() if x==time]
            ypos=spread(e,lo+.13,hi-.13,(hi-lo)*(.07 if time==24 else .053))
            for name,y in e:
                tx={0:2.1,16:18.1,24:26.4}[time]
                ax.annotate(name,xy=(time,y),xytext=(tx,ypos[name]),fontweight='bold',fontsize=8.3,
                    color='#2C3C4A',ha='left',va='center',annotation_clip=False,zorder=10,
                    bbox=dict(boxstyle='round,pad=.13',fc='white',ec='none',alpha=.91),
                    arrowprops=dict(arrowstyle='-',linestyle=(0,(2,2)),color='#697681',lw=.75,shrinkA=1,shrinkB=2))
        counts=subset.groupby('week').size()
        ax.set_title(f'{group}\nn={counts.get(0,0)} / {counts.get(16,0)} / {counts.get(24,0)}',color=GROUP_COLORS[group],fontsize=11.5,fontweight='bold')
        ax.set_xlim(-1.5,32);ax.set_ylim(lo,hi);ax.set_xticks([0,16,24]);ax.set_xlabel('Неделя');ax.grid(axis='y',color='#E5EAEF')
        ax.spines[['top','right']].set_visible(False);ax.tick_params(labelsize=8)
    axes[0].set_ylabel('Доменный балл (центр week 0 = 0)',color=GRAY)
    fig.suptitle(f'{domain} · '+{'D1':'OFT / rear','D2':'Learning / Endurance','D3':'NOR1 + NOR2'}[domain]+' · номер у последнего измерения',fontsize=13,fontweight='bold',color=INK,y=.998)
    fig.subplots_adjust(left=.065,right=.99,top=.80,bottom=.12,wspace=.13)
    path=OUT/f'{domain}_trajectories.png';fig.savefig(path,dpi=155,facecolor='white');plt.close(fig);return path


def plot_mean(df):
    fig,axes=plt.subplots(1,3,figsize=(15.7,3.75),dpi=155)
    for ax,dom in zip(axes,DOMAINS):
        for group in ('PBS','LPS'):
            v=df[df.group==group].groupby('week')[f'{dom}_score'].agg(['mean','sem'])
            x=v.index.to_numpy();y=v['mean'].to_numpy();e=v['sem'].fillna(0).to_numpy()
            ax.plot(x,y,marker='o',lw=2,ms=5,label=group,color=GROUP_COLORS[group]);ax.fill_between(x,y-e,y+e,alpha=.15,color=GROUP_COLORS[group])
        ax.axhline(0,color='#AAB7BF',ls=':',lw=.9);ax.set_xticks([0,16,24]);ax.set_xlabel('Неделя');ax.set_title(dom,color=INK,fontweight='bold');ax.grid(axis='y',color='#E5EAEE');ax.spines[['top','right']].set_visible(False)
        if dom=='D1':ax.set_ylabel('Среднее ± SE')
    axes[-1].legend(frameon=False,fontsize=9)
    fig.suptitle('Средние траектории · PBS / LPS (описательно)',color=INK,fontweight='bold',fontsize=13)
    fig.tight_layout(rect=[0,.01,1,.93]);p=OUT/'mean_trajectories.png';fig.savefig(p,dpi=155);plt.close(fig);return p



def plot_errors(m):
    fig,axes=plt.subplots(3,4,figsize=(17.3,6.75),dpi=150,sharey=True)
    slugs=['ridge','bayes','huber']
    for i,(algo,slug) in enumerate(zip(ALGOS,slugs)):
        for j,(title,_,_,_,_) in enumerate(INTERVALS):
            ax=axes[i,j];rr=m[m.transition==title].set_index('domain').loc[list(DOMAINS)]
            ix=np.arange(3);plain=rr[f'{slug}_plain_mae'].to_numpy();bag=rr[f'{slug}_boot100_mae'].to_numpy()
            ax.bar(ix-.17,plain,width=.32,color='#A2AFB7')
            ax.bar(ix+.17,bag,width=.32,color=COLORS[algo])
            for k,(v,w) in enumerate(zip(plain,bag)):
                if w < v-1e-8:ax.text(k+.17,w+.04,'▼',fontsize=8,ha='center',color=INK)
            ax.set_title(f'{title}\nn={rr.n.iloc[0]} мышей',fontsize=9.9,fontweight='bold',color=INK)
            ax.set_xticks(ix,list(DOMAINS));ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True);ax.spines[['top','right']].set_visible(False)
            if j==0:ax.set_ylabel(f'{algo}\nMAE ↓',color=INK,fontweight='bold')
    fig.legend(handles=[Patch(fc='#A2AFB7',label='Обычная модель'),Patch(fc='#427D9C',label='Bootstrap × 100')],
               loc='upper right',bbox_to_anchor=(.94,.98),ncol=2,frameon=False,fontsize=9)
    fig.suptitle('Bootstrap · MAE обычной модели и усреднения 100 выборок',fontweight='bold',color=INK,fontsize=14,y=.994)
    fig.subplots_adjust(left=.074,right=.985,top=.88,bottom=.07,wspace=.25,hspace=.47)
    path=OUT/'mae_bootstrap_vs_direct.png';fig.savefig(path,dpi=155);plt.close(fig);return path


def plot_winners(m):
    fig,ax=plt.subplots(figsize=(15.7,3.8),dpi=155);ax.set_xlim(-.9,4.06);ax.set_ylim(-.4,3.04);ax.axis('off')
    bgs={'Без изменений':'#F3E3D7','Среднее':'#E6EBEF'}
    for a in ALGOS:
        bgs[a+'|direct']='#ECF0F3';bgs[a+'|bootstrap100']='#DCECF1' if a!='Huber' else '#DBEEE6'
    for j,(title,_,_,_,_) in enumerate(INTERVALS):ax.text(j+.5,2.78,title,ha='center',color=INK,fontweight='bold',fontsize=10.5)
    for i,dom in enumerate(DOMAINS):
        y=2.22-i*.91;ax.text(-.05,y,dom+' — '+{'D1':'OFT','D2':'Learning','D3':'NOR1/2'}[dom],ha='right',va='center',color=INK,fontweight='bold')
        for j,(title,_,_,_,_) in enumerate(INTERVALS):
            row=m[(m.transition==title)&(m.domain==dom)].iloc[0];w=row.best_method
            label=w.replace('|direct','').replace('|bootstrap100',' + boot')
            ax.add_patch(Rectangle((j+.03,y-.35),.94,.73,fc=bgs[w],ec='white',lw=2))
            ax.text(j+.50,y+.11,label,ha='center',va='center',color=INK,fontsize=9.2,fontweight='bold')
            ax.text(j+.50,y-.18,f'MAE {row.best_mae:.3f}',ha='center',va='center',color=GRAY,fontsize=9)
    fig.tight_layout(pad=.4);path=OUT/'best_method_by_domain.png';fig.savefig(path,dpi=155);plt.close(fig);return path


def plot_scatter(m,p):
    fig,axes=plt.subplots(4,3,figsize=(17.6,9.5),dpi=160)
    for i,(title,_,_,_,_) in enumerate(INTERVALS):
        for j,dom in enumerate(DOMAINS):
            ax=axes[i,j];sub=p[(p.transition==title)&(p.domain==dom)]
            row=m[(m.transition==title)&(m.domain==dom)].iloc[0]
            x=sub.actual.to_numpy();y=sub['Ridge|bootstrap100'].to_numpy();lo=min(x.min(),y.min());hi=max(x.max(),y.max());pad=(hi-lo)*.12+.1
            ax.plot([lo-pad,hi+pad],[lo-pad,hi+pad],ls='--',lw=1,color='#9FAAB3')
            ax.scatter(x,y,s=34,color=COLORS['Ridge'],edgecolors='white',linewidths=.5)
            ax.set_xlim(lo-pad,hi+pad);ax.set_ylim(lo-pad,hi+pad)
            ax.spines[['top','right']].set_visible(False);ax.grid(alpha=.17)
            ax.set_title(f'{dom} · MAE {row.ridge_boot100_mae:.3f}',fontsize=11,fontweight='bold',color=INK)
            if j==0:ax.set_ylabel(title+'\nПРОГНОЗ',color=GRAY)
            if i==3:ax.set_xlabel('Фактический балл',color=GRAY)
    fig.suptitle('Ridge + bootstrap · факт против прогноза · пунктир — идеальный прогноз',color=INK,fontweight='bold',fontsize=14,y=.995)
    fig.subplots_adjust(left=.09,right=.985,top=.92,bottom=.07,hspace=.52,wspace=.26)
    path=OUT/'ridge_bootstrap_actual_vs_prediction.png';fig.savefig(path,dpi=155);plt.close(fig);return path


def plot_reps(m):
    fig,axes=plt.subplots(1,3,figsize=(15.7,3.55),dpi=155)
    for ax,(alg,slug) in zip(axes,zip(ALGOS,['ridge','bayes','huber'])):
        xx=[0,25,50,100];yy=[float(m[f'{slug}_{v}_mae'].mean()) for v in ['plain','boot25','boot50','boot100']]
        ax.plot(xx,yy,'o-',color=COLORS[alg],lw=1.8,ms=6)
        ax.set_xticks(xx);ax.grid(alpha=.2)
        ax.set_title(alg,color=INK,fontweight='bold');ax.set_xlabel('Число bootstrap-повторов');ax.set_ylabel('Средняя MAE 12 задач')
        ax.spines[['top','right']].set_visible(False)
    fig.suptitle('Стабилизация результата при 25 / 50 / 100 выборках',fontsize=13,color=INK,fontweight='bold')
    fig.tight_layout(rect=[0,.02,1,.92]);p=OUT/'bootstrap_repeats_sensitivity.png';fig.savefig(p,dpi=155);plt.close(fig);return p

def rgb(hex):return RGBColor(*[int(hex[i:i+2],16) for i in (1,3,5)])

def para(doc,text,bold=False,fontsize=9,color=None,after=5):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.03
    run=p.add_run(text);run.bold=bold;run.font.name='Calibri';run.font.size=Pt(fontsize)
    if color:run.font.color.rgb=rgb(color)
    return p

def heading(doc,s):
    p=para(doc,s,True,17,INK,5);pr=p._p.get_or_add_pPr();bo=OxmlElement('w:pBdr');bot=OxmlElement('w:bottom');bot.set(qn('w:val'),'single');bot.set(qn('w:sz'),'9');bot.set(qn('w:color'),'5484A2');bo.append(bot);pr.append(bo)

def photo(doc,p,width=10.45):
    q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3);q.paragraph_format.space_before=Pt(0)
    q.add_run().add_picture(str(p),width=Inches(width))


def make_doc(df,m,p):
    traj={dom:plot_traj(df,dom) for dom in DOMAINS}
    means=plot_mean(df);errors=plot_errors(m);winners=plot_winners(m);scatter=plot_scatter(m,p);reps=plot_reps(m)
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
    sec.left_margin=Inches(.44);sec.right_margin=Inches(.44);sec.top_margin=Inches(.38);sec.bottom_margin=Inches(.40);sec.footer_distance=Inches(.16)
    doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(9)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;rr=foot.add_run('24.10.2026 · eNRI / 3 домена · bootstrap');rr.font.size=Pt(8);rr.font.color.rgb=rgb('#7D8990')
    heading(doc,'Динамика трёх доменов · PBS и LPS')
    sub=df[df.group.isin(['PBS','LPS'])]
    para(doc,f'{sub.mouse_id.nunique()} мышей · {len(sub)} живых состояний · 0 / 16 / 24 недели. Сплошные линии — мыши; тёмный пунктир — линии к номерам.',fontsize=8.8,color=GRAY)
    para(doc,'D1 · двигательная активность — OFT / rear',True,11,INK,2);photo(doc,traj['D1'],10.35)
    para(doc,'D2 · обучение и выносливость — Learning / Endurance',True,11,INK,2);photo(doc,traj['D2'],10.35)
    para(doc,'Вывод: индивидуальные изменения разнонаправлены; 0 означает центр исходной шкалы.',fontsize=8.5,color=GRAY)
    doc.add_page_break()
    heading(doc,'Динамика доменов · продолжение')
    para(doc,'D3 · исследовательское поведение — NOR1 + NOR2',True,11,INK,2);photo(doc,traj['D3'],10.35)
    para(doc,'Средние баллы D1–D3 · PBS / LPS',True,11,INK,2);photo(doc,means,10.25)
    para(doc,'Bootstrap: 100 выборок мышей С ВОЗВРАЩЕНИЕМ только из обучающей части; прогнозы 100 отдельных моделей усредняются. Исключённая мышь не участвует в обучении. Не создаются новые независимые животные.',fontsize=8.4,color=GRAY,after=2)
    doc.add_page_break()
    heading(doc,'Точность прогноза · обычная модель / bootstrap')
    para(doc,'12 задач, n=6–11. Leave-one-mouse-out. Ridge, Bayesian Ridge, Huber: 100 bootstrap-выборок для каждого исключения мыши; MAE меньше = лучше.',fontsize=8.8,color=GRAY,after=3)
    photo(doc,errors,9.55)
    para(doc,'Лучший способ · три регрессии с bootstrap/без него и два простых прогноза',True,10.3,INK,2)
    photo(doc,winners,9.8)
    wins={alg:int((m[f'{slug}_boot_gain']>0).sum()) for alg,slug in zip(ALGOS,['ridge','bayes','huber'])}
    best=int(m.bootstrap_best_of_8.sum())
    para(doc,f'Итог: bootstrap уменьшил MAE в {wins["Ridge"]}/12 задач Ridge, {wins["Bayesian Ridge"]}/12 Bayesian Ridge, {wins["Huber"]}/12 Huber. Bootstrap-модель лучшая из восьми в {best}/12.',True,8.8,INK,after=2)
    doc.add_page_break()
    heading(doc,'Факт против прогноза · 12 проверок Ridge + bootstrap')
    para(doc,'Ряды: PBS 0→16 (n=9), PBS 16→24 (n=6), LPS 16→24 (n=11), PBS 0→24 (n=6). Столбцы D1–D3; пунктир — идеальный прогноз.',fontsize=8.8,color=GRAY)
    photo(doc,scatter,10.5)
    para(doc,'Вывод: bootstrap оценивает зависимость прогноза от обучающих мышей и может снизить вариабельность алгоритма, но не увеличивает реальный размер выборки.',True,8.4,INK,2)
    para(doc,'Ограничения: EFA и исходная импутация выполнены до CV; проверена только регрессия при фиксированных доменных баллах. Выборка ограничена мышами с обеими точками; n=6–11, нет внешней валидации.',fontsize=7.7,color=GRAY,after=0)
    file=OUT/'Отчет_Bootstrap.docx';doc.save(file);return file


def main():
    df=pd.read_csv(DATA,dtype={'mouse_id':str});assert len(df)==133
    m,p=calculate(df)
    doc=make_doc(df,m,p)
    wins={alg:int((m[f'{slug}_boot_gain']>0).sum()) for alg,slug in zip(ALGOS,['ridge','bayes','huber'])}
    means={alg:{'plain':float(m[f'{slug}_plain_mae'].mean()),'bootstrap100':float(m[f'{slug}_boot100_mae'].mean())} for alg,slug in zip(ALGOS,['ridge','bayes','huber'])}
    summary={'n_tasks':12,'repetitions':REPS,'repeats_sensitivity':CHECK_REPS,'rng_seed':SEED,'bootstrap_beats_plain':wins,
             'bootstrap_best_of_eight':int(m.bootstrap_best_of_8.sum()),'mean_mae_by_algorithm':means,
             'mean_unique_per_bootstrap_draw':float(m.mean_unique_training_mice.mean()),
             'transition_n':{name:n for name,_,_,_,n in INTERVALS},'historical_ridge_verified':True}
    (OUT/'SUMMARY.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    lines=['# Bootstrap aggregation on 3 rank-EFA domains / PBS and LPS','',
         'Within each LOOCV training fold, 100 bootstrap resamples of real training mouse-pairs (n−1 each, with replacement). Independent models fit per draw; predictions averaged. NO synthetic mice. Original heldout mouse never used in fitting/scaling.','',
         '| Transition | Domain | n | Ridge | Ridge boot | Bayes | Bayes boot | Huber | Huber boot | Persistence | Training mean | Best of 8 |',
         '|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---|']
    for r in m.itertuples():
        lines.append(f'| {r.transition} | {r.domain} | {r.n} | {r.ridge_plain_mae:.3f} | {r.ridge_boot100_mae:.3f} | {r.bayes_plain_mae:.3f} | {r.bayes_boot100_mae:.3f} | {r.huber_plain_mae:.3f} | {r.huber_boot100_mae:.3f} | {r.unchanged_mae:.3f} | {r.mean_mae:.3f} | {r.best_method} |')
    lines += ['',f'Bootstrap beats direct in {wins}; bootstrap best of eight in {summary["bootstrap_best_of_eight"]}/12.',
              f'Mean MAE: {means}. Mean number unique training mice in a bootstrap draw: {summary["mean_unique_per_bootstrap_draw"]:.2f}.',
              'Sensitivity 25/50/100 bootstraps is only a Monte Carlo convergence check, not hyperparameter optimization. No claims of reliable uncertainty intervals.',
              'Input EFA/imputation precomputed outside outer LOOCV. Survivors-only paired samples; n=6–11, no external test.','',
              'Word report: results/Отчет_Bootstrap.docx']
    (OUT/'REPORT.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    print(m[['transition','domain','n','ridge_plain_mae','ridge_boot100_mae','bayes_plain_mae','bayes_boot100_mae','huber_plain_mae','huber_boot100_mae','best_method']].round(3).to_string(index=False))
    print('FILE',doc,doc.stat().st_size)

if __name__=='__main__':main()
